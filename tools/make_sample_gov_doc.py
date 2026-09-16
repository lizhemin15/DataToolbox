#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成示例公文 Word：省-市-区-县 四级标题 + 每县 4 段（人口/经济/工业/教育）文字描述。"""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "apps", "data-ontology", "example_files", "区市县经济社会发展情况通报.docx")

PROV = "江源省"
TREE = [
    ("云台市", [
        ("城东区", [
            ("平安县", dict(pop="28.6", pop_add="0.4", urb="43.4", urb_add="1.2", gdp="156.3", gdp_yoy="5.8",
                            fis="9.2", retail="62.4", ind_n="87", ind_v="41.2", ind_yoy="6.5",
                            sch="126", stu="4.8", rate="99.2")),
            ("长乐县", dict(pop="21.3", pop_add="0.2", urb="38.6", urb_add="0.9", gdp="118.7", gdp_yoy="5.1",
                            fis="6.8", retail="45.1", ind_n="64", ind_v="29.5", ind_yoy="5.7",
                            sch="98", stu="3.6", rate="98.8")),
        ]),
        ("城西区", [
            ("清溪县", dict(pop="33.9", pop_add="0.6", urb="47.2", urb_add="1.5", gdp="201.5", gdp_yoy="6.2",
                            fis="12.4", retail="78.9", ind_n="112", ind_v="57.8", ind_yoy="7.1",
                            sch="152", stu="6.1", rate="99.4")),
            ("白沙县", dict(pop="17.8", pop_add="0.1", urb="34.9", urb_add="0.7", gdp="92.4", gdp_yoy="4.8",
                            fis="5.3", retail="36.7", ind_n="48", ind_v="21.6", ind_yoy="5.2",
                            sch="76", stu="2.7", rate="98.5")),
        ]),
    ]),
    ("临江市", [
        ("江北区", [
            ("新丰县", dict(pop="25.4", pop_add="0.3", urb="41.1", urb_add="1.1", gdp="143.2", gdp_yoy="5.6",
                            fis="8.6", retail="57.3", ind_n="79", ind_v="37.4", ind_yoy="6.0",
                            sch="114", stu="4.2", rate="99.0")),
            ("石城县", dict(pop="19.6", pop_add="0.2", urb="36.3", urb_add="0.8", gdp="104.9", gdp_yoy="5.0",
                            fis="6.1", retail="41.2", ind_n="56", ind_v="25.3", ind_yoy="5.4",
                            sch="88", stu="3.1", rate="98.6")),
        ]),
        ("江南区", [
            ("龙泉县", dict(pop="31.2", pop_add="0.5", urb="45.8", urb_add="1.4", gdp="186.7", gdp_yoy="6.0",
                            fis="11.3", retail="71.6", ind_n="103", ind_v="52.4", ind_yoy="6.8",
                            sch="141", stu="5.5", rate="99.3")),
            ("永安县", dict(pop="15.9", pop_add="0.1", urb="32.7", urb_add="0.6", gdp="83.6", gdp_yoy="4.6",
                            fis="4.8", retail="32.5", ind_n="42", ind_v="18.9", ind_yoy="4.9",
                            sch="69", stu="2.4", rate="98.3")),
        ]),
    ]),
]


def paras(county, d):
    return [
        f"{county}常住人口 {d['pop']} 万人，比上年末增加 {d['pop_add']} 万人；城镇化率 {d['urb']}%，较上年提高 {d['urb_add']} 个百分点。",
        f"{county}实现地区生产总值 {d['gdp']} 亿元，同比增长 {d['gdp_yoy']}%；一般公共预算收入 {d['fis']} 亿元，社会消费品零售总额 {d['retail']} 亿元。",
        f"{county}规模以上工业企业 {d['ind_n']} 家，完成工业增加值 {d['ind_v']} 亿元，同比增长 {d['ind_yoy']}%。",
        f"{county}现有各级各类学校 {d['sch']} 所，在校学生 {d['stu']} 万人，九年义务教育巩固率 {d['rate']}%。",
    ]


doc = Document()
style = doc.styles['Normal']
style.font.name = '仿宋_GB2312'
style.font.size = Pt(14)

t = doc.add_paragraph('全省区市县经济社会发展情况通报')
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in t.runs:
    r.bold = True
    r.font.size = Pt(18)

doc.add_paragraph('各市（区）、县：')
doc.add_paragraph('现将全省各县（区）人口、经济、工业、教育等方面情况通报如下。')

doc.add_paragraph(f'一、{PROV}')
for city, districts in TREE:
    doc.add_paragraph(f'（一）{city}')
    for i, (dist, counties) in enumerate(districts, 1):
        doc.add_paragraph(f'{i}. {dist}')
        for j, (county, data) in enumerate(counties, 1):
            doc.add_paragraph(f'（{j}）{county}')
            for p in paras(county, data):
                doc.add_paragraph(p)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print('已生成:', OUT)
print('段落数:', len(doc.paragraphs))
