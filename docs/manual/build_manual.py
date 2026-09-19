#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成《数据工具箱使用说明书》（公文格式排版，图文并茂）。

用法：
    python3 docs/manual/build_manual.py
输出：
    docs/使用说明书.docx

说明：
- 排版遵循党政机关公文格式（GB/T 9704-2012）常用约定：
  页边距 上37mm/下35mm/左28mm/右26mm；正文三号仿宋_GB2312，行距固定值28磅，首行缩进2字符；
  标题二号方正小标宋简体；一级标题三号黑体、二级三号楷体_GB2312 加粗、三级三号仿宋_GB2312 加粗。
- 截图由 docs/manual/take_screenshots.js 生成，放在 docs/manual/screenshots/。
- 功能有更新时：更新下面 CONTENT 里的文字 / 表格，必要时重跑截图脚本，然后重新生成并提交。
"""

import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, 'screenshots')
OUT = os.path.join(os.path.dirname(HERE), '使用说明书.docx')

VERSION = os.environ.get('DT_MANUAL_VERSION', 'v2026.9.19115018')
DATE_CN = os.environ.get('DT_MANUAL_DATE', '二〇二六年九月')

# 字体
F_FS = '仿宋_GB2312'      # 正文
F_HT = '黑体'             # 一级标题
F_KT = '楷体_GB2312'      # 二级标题 / 图注
F_XBS = '方正小标宋简体'   # 大标题
F_ST = '宋体'             # 页码 / 表格

SZ_TITLE = 22   # 二号
SZ_BODY = 16    # 三号
SZ_SMALL = 12   # 小四
SZ_CAP = 10.5   # 五号

LINE = 28       # 正文行距（磅）


# ---------------------------- 公文排版基础工具 ----------------------------

def set_font(run, east=F_FS, ascii_f='Times New Roman', size=SZ_BODY, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = ascii_f
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.insert(0, rf)
    rf.set(qn('w:ascii'), ascii_f)
    rf.set(qn('w:hAnsi'), ascii_f)
    rf.set(qn('w:eastAsia'), east)


def set_indent_chars(par, first=2, left=0):
    """按“字符”设置缩进（公文要求首行缩进 2 字符）"""
    pPr = par._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    if first:
        ind.set(qn('w:firstLineChars'), str(int(first * 100)))
        ind.set(qn('w:firstLine'), '0')
    if left:
        ind.set(qn('w:leftChars'), str(int(left * 100)))


def set_outline(par, level):
    """给自定义标题标记大纲级别，便于 Word 生成目录"""
    pPr = par._p.get_or_add_pPr()
    ol = OxmlElement('w:outlineLvl')
    ol.set(qn('w:val'), str(level))
    pPr.append(ol)


def body(doc, text, indent=2, size=SZ_BODY, font=F_FS, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         line=LINE, space_after=0, color=None):
    par = doc.add_paragraph()
    par.alignment = align
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line)
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    set_indent_chars(par, first=indent)
    run = par.add_run(text)
    set_font(run, east=font, size=size, bold=bold)
    if color:
        run.font.color.rgb = color
    return par


def title(doc, text, size=SZ_TITLE, font=F_XBS, space_after=18, space_before=0):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(size + 12)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    run = par.add_run(text)
    set_font(run, east=font, size=size, bold=False)
    return par


HEADINGS = []          # [(level, text)] 供目录使用
TOC_ANCHORS = []       # 需要在目录里列出的标题（一级标题 + 附录）


def h1(doc, text):
    par = body(doc, text, indent=2, font=F_HT, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)
    set_outline(par, 0)
    TOC_ANCHORS.append(text)
    return par


def h2(doc, text):
    par = body(doc, text, indent=2, font=F_KT, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_outline(par, 1)
    return par


def toc_anchor(doc, text, size=None, font=None, space_after=10):
    """附录一类的标题：也进目录"""
    par = title(doc, text, size=size or (SZ_BODY + 2), font=font or F_HT, space_after=space_after,
                space_before=0)
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_indent_chars(par, first=2)
    set_outline(par, 0)
    TOC_ANCHORS.append(text)
    return par


def h3(doc, text):
    par = body(doc, text, indent=2, font=F_FS, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_outline(par, 2)
    return par


def bullet(doc, text, size=SZ_BODY, font=F_FS):
    return body(doc, text, indent=2, size=size, font=font)


def image(doc, filename, caption=None, width_cm=15.0):
    path = os.path.join(SHOTS, filename)
    if not os.path.exists(path):
        body(doc, '（图缺失：%s）' % filename, indent=2, color=RGBColor(0xC0, 0x00, 0x00))
        return
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = par.add_run()
    run.add_picture(path, width=Cm(width_cm))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r = cap.add_run(caption)
        set_font(r, east=F_KT, size=SZ_CAP)
        r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)


def table(doc, rows, widths=None, header=True, size=SZ_CAP, font=F_ST):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ''
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = par.add_run(str(cell_text))
            set_font(run, east=font, size=size, bold=(header and i == 0))
    if widths:
        for j, w in enumerate(widths):
            for i in range(len(rows)):
                t.cell(i, j).width = Cm(w)
    # 表格后空一行
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return t


def code_block(doc, lines, size=SZ_SMALL):
    for ln in lines:
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Cm(0.75)
        run = par.add_run(ln if ln else ' ')
        set_font(run, east='Consolas', ascii_f='Consolas', size=size)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)
    sp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def setup_document():
    doc = Document()

    # 默认样式（正文）
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(SZ_BODY)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), F_FS)

    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(3.7)
        sec.bottom_margin = Cm(3.5)
        sec.left_margin = Cm(2.8)
        sec.right_margin = Cm(2.6)
        sec.header_distance = Cm(1.5)
        sec.footer_distance = Cm(2.0)
        # 页脚页码（四号宋体，居中）
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run()
        set_font(run, east=F_ST, size=14)
        fld_begin = OxmlElement('w:fldChar')
        fld_begin.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = 'PAGE'
        fld_end = OxmlElement('w:fldChar')
        fld_end.set(qn('w:fldCharType'), 'end')
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)
    return doc


def add_toc(doc):
    par = doc.add_paragraph()
    run = par.add_run()
    set_font(run, east=F_ST, size=SZ_BODY)
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    begin.set(qn('w:dirty'), 'true')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = r'TOC \o "1-3" \h \z \u'
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = '（目录域：在 Word 中按 Ctrl+A 后 F9，或右键“更新域”即可生成带页码的目录）'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for e in (begin, instr, sep, t, end):
        run._r.append(e)


# ---------------------------- 说明书正文内容 ----------------------------

def build_cover(doc, toc_items=None, pages=None):
    title(doc, '数据工具箱使用说明书', size=SZ_TITLE, space_before=60, space_after=24)
    body(doc, '（版本 %s）' % VERSION, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER, font=F_KT, size=SZ_BODY, space_after=6)
    body(doc, '数据工具箱项目组', indent=0, align=WD_ALIGN_PARAGRAPH.CENTER, font=F_KT, size=SZ_BODY, space_after=6)
    body(doc, DATE_CN, indent=0, align=WD_ALIGN_PARAGRAPH.CENTER, font=F_KT, size=SZ_BODY, space_after=6)
    body(doc, '', indent=0)
    body(doc, '本说明书随版本发布同步更新，如有疑问请以系统内实际界面为准。',
         indent=0, align=WD_ALIGN_PARAGRAPH.CENTER, font=F_KT, size=SZ_CAP)
    doc.add_page_break()

    title(doc, '目　　录', size=SZ_BODY + 6, font=F_HT, space_after=12)
    if toc_items:
        build_toc(doc, toc_items, pages or {})
    doc.add_page_break()


def build_intro(doc):
    toc_anchor(doc, '一、系统概述', size=SZ_TITLE - 6)

    h2(doc, '（一）系统定位')
    body(doc, '数据工具箱（DataToolbox）是一套面向业务人员与技术人员的一体化数据处理与治理平台，'
              '采用“单文件服务 + 浏览器界面”的轻量形态部署，无需安装数据库客户端，'
              '打开浏览器即可完成数据接入、查询、治理、分发与可视化。')
    body(doc, '平台的核心思路是“把日常的数据处理动作沉淀成可复用的任务”：'
              '一次写好处理逻辑，之后既可以手工上传文件执行，也可以定时自动执行，'
              '还可以注册成接口分发给其它系统调用。')

    h2(doc, '（二）功能总览')
    body(doc, '平台顶部导航按模块划分，各模块的主要能力如下表所示。')
    table(doc, [
        ['模块', '主要能力', '常用人群'],
        ['数据库管理', '接入多种数据库、浏览表结构与数据、执行 SQL、结果导出', '技术/分析人员'],
        ['数据治理', '把数据加工逻辑固化为任务（定时/交互），支持文件解析、模板填充、AI 处理、接口分发', '数据工程师'],
        ['本体论抽象', '从表结构抽取本体概念与关系，形成业务语义模型（默认隐藏，需在设置中开启）', '数据架构师'],
        ['数据血缘', '展示表与表、表与任务之间的流转关系（默认隐藏，需在设置中开启）', '数据治理人员'],
        ['接口分发', '管理已注册的对外接口、查看接口文档与在线调试', '集成开发人员'],
        ['Agent服务', '以 MCP 协议向智能体（Agent）暴露平台的数据库与治理能力', 'AI 应用开发者'],
        ['智能助手', '与大模型对话，辅助生成 SQL、生成治理任务、解释数据口径', '业务/技术通用'],
        ['模型管理', '维护可用的模型清单与调用参数（默认隐藏，需在设置中开启）', '平台管理员'],
        ['数据质量审核', '按规则对数据填报情况做审核，输出审核结论与整改建议', '业务主管'],
        ['应用广场', '一键启用预制的看板/表格类应用模板', '业务人员'],
        ['大屏编辑器', '拖拽式拼装数据大屏（默认隐藏，需在设置中开启）', '可视化人员'],
    ], widths=[3.0, 9.5, 3.0])

    h2(doc, '（三）运行环境与兼容性')
    table(doc, [
        ['项目', '要求'],
        ['服务端操作系统', 'Linux（amd64/arm64）、Windows（amd64/arm64）'],
        ['服务端资源', '推荐 2 核 4GB 及以上；纯前端执行的任务对服务端压力极小'],
        ['浏览器', 'Chrome / Edge 100 及以上（推荐最新版）；不支持 IE'],
        ['端口', '默认 8080，可在 server.config.json 或启动参数中修改'],
        ['数据存储', '默认使用内置文件型存储（data/data-store.json），无需额外安装数据库'],
        ['可选外部依赖', '如需 AI 能力，需配置一个 OpenAI 兼容的大模型服务地址与密钥'],
    ], widths=[4.0, 11.5])


def build_install(doc):
    toc_anchor(doc, '二、安装与部署', size=SZ_TITLE - 6)

    h2(doc, '（一）安装包内容')
    body(doc, '各平台发布包（ZIP / tar.gz）解压后目录结构基本一致，主要文件如下。')
    table(doc, [
        ['文件 / 目录', '说明'],
        ['datatoolbox-server(.exe)', '服务端主程序，双击或命令行启动即可'],
        ['install.sh / update.sh', 'Linux 交互式安装脚本与更新脚本'],
        ['server.config.json', '服务配置（端口、监听地址）'],
        ['README_DEPLOY.md', '部署说明（命令行速查）'],
        ['使用说明书.docx', '本说明书'],
        ['index.html、js/、css/、lib/', '前端页面与静态资源'],
        ['apps/、templates/、components/', '子应用、看板模板与预制组件'],
        ['scripts/、examples/、assets/', '治理脚本、示例文件与静态资源'],
        ['data/', '数据目录（首次启动自动生成 data-store.json）'],
    ], widths=[6.0, 9.5])

    h2(doc, '（二）Linux 安装（推荐使用脚本）')
    body(doc, '第一步，解压发布包并进入目录：')
    code_block(doc, ['tar -xzf datatoolbox-linux-amd64.tar.gz', 'cd datatoolbox-linux-amd64'])
    body(doc, '第二步，执行安装脚本，按提示回答：')
    code_block(doc, ['chmod +x install.sh', './install.sh'])
    body(doc, '脚本会依次询问安装目录（root 默认 /opt/datatoolbox，普通用户默认 ~/datatoolbox）、'
              '监听端口（默认 8080）、是否配置 systemd 服务与开机自启，并在完成后做端口连通性验证。')
    body(doc, '第三步，验证服务：')
    code_block(doc, ['systemctl status datatoolbox', 'curl -I http://127.0.0.1:8080/'])
    body(doc, '浏览器访问 http://服务器地址:8080，使用管理员账号登录即可（初始账号 admin / admin1234，'
              '首次登录后请立即修改密码）。')

    h2(doc, '（三）Windows 安装')
    body(doc, '下载对应架构的 zip 包解压，双击 datatoolbox-server.exe 即可启动；'
              '如需开机自启，可将其注册为 Windows 服务（可用 nssm 等工具），或在“任务计划程序”中添加开机任务。'
              '启动后同样通过浏览器访问 http://127.0.0.1:8080。')

    h2(doc, '（四）手动部署与启动参数')
    body(doc, '不依赖脚本时，编辑 server.config.json 后直接运行主程序：')
    code_block(doc, ['{', '  "port": 8080,', '  "host": "0.0.0.0"', '}'])
    code_block(doc, ['chmod +x datatoolbox-server', './datatoolbox-server', './datatoolbox-server -port 3000   # 命令行参数优先'])
    body(doc, '服务日志默认输出到标准输出；使用 install.sh 配置 systemd 时，会写入日志目录（'
              'root 且安装到 /opt/datatoolbox 时为 /var/log/datatoolbox/server.log，否则为安装目录下 logs/server.log）。')

    h2(doc, '（五）升级更新')
    body(doc, '升级只需替换程序与前端静态文件，数据保存在 data/ 目录不受影响。推荐使用更新脚本：')
    code_block(doc, ['./update.sh'])
    body(doc, '手动升级时：停止服务 → 备份 data/ 与 server.config.json → 用新版本覆盖程序与前端文件 → '
              '启动服务 → 浏览器强制刷新（Ctrl+F5）确认版本号。')

    h2(doc, '（六）部署常见问题')
    table(doc, [
        ['现象', '排查方向'],
        ['页面打不开', '确认服务进程存活、端口未被占用（netstat -lntp | grep 8080）、防火墙/安全组已放行'],
        ['登录后空白', '清理浏览器缓存并强制刷新；确认 index.html 与 js/ 目录为同一版本'],
        ['上传大文件失败', '确认磁盘空间、反向代理（Nginx）的 client_max_body_size 配置'],
        ['AI 功能不可用', '在“设置 → LLM 配置”中检查 URL / API Key / 模型名是否正确，可点“自动获取模型列表”验证连通性'],
        ['定时任务不执行', '确认任务已启用、服务持续运行；查看任务执行日志'],
    ], widths=[4.5, 11.0])


def build_login_and_nav(doc):
    toc_anchor(doc, '三、登录与界面总览', size=SZ_TITLE - 6)

    h2(doc, '（一）登录')
    body(doc, '在浏览器地址栏输入服务地址（如 http://127.0.0.1:8080）后进入登录页，'
              '输入用户名与密码即可进入系统。管理员初始账号为 admin / admin1234，首次登录后请及时修改密码。')
    image(doc, '01-login.png', '图 1　系统登录界面')

    h2(doc, '（二）顶部导航与页面结构')
    body(doc, '登录后的主界面自上而下分为三个区域：顶部为系统名称、功能标签页与用户操作区；'
              '中部为各模块的工作区；底部与右侧为操作反馈（提示条、执行日志、结果表格等）。'
              '点击标签页即可切换模块，当前模块高亮显示。')
    image(doc, '02-database.png', '图 2　主界面与顶部导航（以“数据库管理”为例）')

    h2(doc, '（三）标签页的显示与排序')
    body(doc, '系统默认只显示常用模块，其余模块可按需开启。点击右上角“设置”按钮，'
              '在“标签页设置”区域勾选需要显示的模块、调整顺序或重命名，保存后立即生效。'
              '该设置按用户保存，不同账号互不影响。')
    image(doc, '20-tab-settings.png', '图 3　设置界面中的标签页显示与排序')

    h2(doc, '（四）通用交互约定')
    bullet(doc, '1. 操作反馈：成功/失败均以右上角提示条呈现，失败时会给出原因。')
    bullet(doc, '2. 长任务：耗时操作（如 AI 生成、批量文件处理）会显示进度或流式输出，可等待完成后再操作其它模块。')
    bullet(doc, '3. 结果表格：查询或解析结果以表格形式展示，支持横向滚动，重要结果可导出为 Excel。')
    bullet(doc, '4. 浏览器建议：请勿使用 IE；版本升级后建议按 Ctrl+F5 强制刷新，避免加载到旧版静态资源。')


def build_database(doc):
    toc_anchor(doc, '四、数据库管理', size=SZ_TITLE - 6)

    h2(doc, '（一）添加数据库连接')
    body(doc, '在“数据库管理”模块点击“添加数据库”，填写连接信息并保存：')
    table(doc, [
        ['字段', '说明'],
        ['名称', '自定义的展示名称，建议体现业务含义'],
        ['类型', 'MySQL、PostgreSQL、SQL Server、Oracle、达梦、人大金仓等'],
        ['主机 / 端口', '数据库服务地址与端口'],
        ['库名 / 用户名 / 密码', '用于连接的库与账号；建议使用只读账号做查询类接入'],
        ['备注', '可选的用途说明，便于多人协作辨识'],
    ], widths=[4.5, 11.0])
    body(doc, '保存后可点击“测试连接”验证连通性；连接信息保存在服务端数据库中，'
              '请注意账号权限最小化，避免把高权限账号交给非管理员使用。')

    h2(doc, '（二）表结构浏览')
    body(doc, '选中数据库后，左侧为表清单，右侧展示所选表的字段结构（字段名、类型、注释等）。'
              '支持按表名搜索、查看建表语句与字段口径，是编写数据治理任务时确认字段的依据。')

    h2(doc, '（三）数据查询与结果导出')
    body(doc, '在 SQL 编辑区输入查询语句并执行，结果以表格返回；'
              '可将结果导出为 Excel，或一键“另存为治理任务”，把查询逻辑固化为可复用任务。')

    h2(doc, '（四）数据安全建议')
    bullet(doc, '1. 生产库接入优先使用只读账号；需要写入时单独申请权限。')
    bullet(doc, '2. 定期检查已保存的连接与密码，人员变动后及时更换。')
    bullet(doc, '3. 导出文件含敏感数据时，请遵循单位的数据分级分类管理要求。')


def build_governance(doc):
    toc_anchor(doc, '五、数据治理', size=SZ_TITLE - 6)
    body(doc, '数据治理是本平台使用频率最高的模块。它的本质是：把“拿到数据 → 清洗加工 → 生成结果/入库”'
              '这套动作写成一个任务，之后可以反复执行。')

    h2(doc, '（一）两类任务')
    table(doc, [
        ['类型', '适用场景', '执行触发方式'],
        ['定时任务', '周期性同步、汇总、抽取（如每天凌晨统计报表）', '按 Cron 表达式自动执行，也可手工点“运行”'],
        ['交互任务', '需要人工上传文件或粘贴文本后处理（如把 Word 公告转成 Excel）', '上传文件/输入文本后点“执行任务”，也可点“运行”'],
    ], widths=[2.5, 9.0, 4.0])

    h2(doc, '（二）任务列表与新建任务')
    body(doc, '进入“数据治理”模块，左侧为任务列表，右侧为所选任务的详情。'
              '点击“新建任务”打开表单，各字段含义如下。')
    image(doc, '05-governance-tasks.png', '图 4　数据治理任务列表')
    table(doc, [
        ['字段', '说明'],
        ['任务名称', '必填，建议体现“对象 + 动作 + 结果”，如“产品汇总Word生成（华东区）”'],
        ['任务类型', '定时任务 / 交互任务，决定执行方式与可配置项'],
        ['任务描述', '给同事看的说明，写清输入是什么、输出是什么'],
        ['关联数据库', '可选。关联后任务脚本中可用 gov.querySQL / gov.executeSQL 操作该库'],
        ['Cron 表达式', '定时任务专用，格式为“分 时 日 月 周”，如 0 2 * * * 表示每天 02:00'],
        ['输入类型', '交互任务专用：文件 / 文本 / 两者都要'],
        ['允许的文件类型', '交互任务专用，如 .docx,.xlsx,.csv'],
        ['多文件模式', '逐文件执行（per_file）或合并为一次执行（single，脚本中用 INPUT_FILES 取全部文件）'],
        ['运行环境', '前端（浏览器内执行）或后端运行器（服务端执行）'],
        ['任务代码', 'JavaScript 逻辑，可调用平台内置的 gov.* API（见附录 A）'],
    ], widths=[3.6, 11.9])

    h2(doc, '（三）任务详情页要素')
    body(doc, '选中任务后，详情页自上而下依次为：任务名称与操作按钮（编辑、运行、分享、复制、AI 新建、删除）、'
              '基本信息（类型、状态、Cron、关联数据库、最后运行时间）、数据输入区（交互任务）、'
              '任务代码折叠区、执行日志区。')
    image(doc, '06-gov-task-detail.png', '图 5　任务详情页（含“复制”“AI 新建”等操作按钮）')

    h2(doc, '（四）运行任务与查看执行结果')
    body(doc, '点击“▶️ 运行”即可执行任务：带输入的任务会自动带上已选择的文件与文本；'
              '执行过程与结果输出在执行日志区，结构化结果会以表格形式预览，'
              '若生成了文件（如 Excel、Word），日志下方会出现下载入口。')
    image(doc, '08-gov-run-output.png', '图 6　执行日志与结果预览')

    h2(doc, '（五）交互任务的输入方式')
    body(doc, '交互任务在“数据输入”区域提供输入：文件可拖拽或点击上传（支持多选），文本类任务提供多行输入框。'
              '多文件模式为“逐文件执行”时，每个文件单独跑一次任务；为“合并执行”时，'
              '全部文件一次性传给脚本，脚本中通过 INPUT_FILES 数组读取。')
    image(doc, '07-gov-input-files.png', '图 7　交互任务的文件上传与输入区域')

    h2(doc, '（六）执行环境：前端与后端')
    table(doc, [
        ['运行环境', '执行位置', '特点与适用'],
        ['前端', '浏览器内执行', '文件不上传服务端，适合含敏感数据的 Word/Excel 解析、模板填充；依赖浏览器性能，超大文件较慢'],
        ['后端', '服务端运行器', '适合定时任务、大数据量处理、需要长期后台执行的场景'],
    ], widths=[2.8, 3.2, 9.5])

    h2(doc, '（七）内置示例任务')
    body(doc, '发布包自带一批可直接使用的示例任务，覆盖 Word/Excel/CSV 解析、模板填充、AI 抽取、'
              '定时检查等典型场景，可作为编写新任务的模板参考。示例任务的输入样例文件可在任务详情页'
              '“示例文件”处直接下载。详见附录 B。')

    h2(doc, '（八）分享链接')
    body(doc, '对需要他人临时使用的任务，可点击“🔗 分享”生成分享链接。'
              '获得链接的人无需登录即可在受限页面提交输入并取回结果，适合一次性文件转换等场景。'
              '不需要时请及时关闭分享，避免链接外泄造成数据暴露。')

    h2(doc, '（九）注册为 API')
    body(doc, '任务可注册为对外接口，供其它系统调用。注册后可在“接口分发”模块查看调用方式与在线调试。'
              '接口调用需要平台颁发的 API Key（在“设置 → API Key”中管理）。')

    h2(doc, '（十）复制任务：一键派生同类任务')
    body(doc, '当需要“同样的逻辑、换一套参数”的任务时，不必从零编写：'
              '打开作为模板的任务，点击“📄 复制”，输入新任务名称即可生成一个副本。'
              '副本会完整继承原任务的代码与配置，但不会继承运行记录、分享链接与 API 注册，状态重置为待运行。')
    body(doc, '典型用法：以“产品汇总Word生成（套用表格模板）”为模板，复制出“产品汇总Word生成（华东区）”，'
              '再进入“编辑”微调输出文件名或列定义。')

    h2(doc, '（十一）AI 新建任务：先问清需求，再生成任务')
    body(doc, '对于“想要一个新任务、但不确定怎么改代码”的场景，可使用“✨ AI 新建”。'
              '它不是让 AI 凭空造一个任务，而是以当前任务为模板，分三步完成：')
    bullet(doc, '第一步（AI 提问）：AI 先阅读模板任务的说明与代码，然后针对这个模板提出 3～6 个问题，'
                '例如“要抽取的字段与模板有何不同”“输出文件名与格式”“是单文件还是多文件”“是否需要调用 AI”。'
                '不同模板问的问题不同，因为这些差异点本身就是任务特异化的地方。')
    bullet(doc, '第二步（填写需求）：逐条回答问题，回答越具体，生成结果越贴合预期。'
                '例如写明“每个产品单独一张表格，套用模板样式”，而不要只写“格式好看一点”。')
    bullet(doc, '第三步（生成并保存）：AI 生成任务名称、说明与完整代码，界面会展示代码预览，'
                '确认无误后点击“保存为新任务”。保存后仍可用“编辑”继续微调，或直接运行验证。')
    image(doc, '09-ai-clone-step1.png', '图 8　AI 新建向导第一步：由 AI 针对模板提问')
    image(doc, '10-ai-clone-questions.png', '图 9　AI 新建向导第二步：填写需求回答')
    image(doc, '11-ai-clone-generated.png', '图 10　AI 新建向导第三步：生成任务代码并预览')

    h2(doc, '（十二）编写任务脚本的注意事项')
    bullet(doc, '1. 先读样例、再动手：优先复制一个功能相近的示例任务改，而不是从空白写起。')
    bullet(doc, '2. 用 gov.log 输出关键中间结果，便于排查；结构化结果用 gov.showTable 展示。')
    bullet(doc, '3. 文件读取/写入遵守同一套 API（gov.readWord / gov.writeExcel / gov.word 等），'
                '不要引入模板里没有用过的第三方库。')
    bullet(doc, '4. 多文件任务注意 INPUT_FILES 与单文件变量 INPUT_FILE 的区别。')
    bullet(doc, '5. 涉及数据库写入的任务，先在测试库验证，避免误写生产数据。')


def build_other_modules(doc):
    toc_anchor(doc, '六、本体论抽象', size=SZ_TITLE - 6)
    body(doc, '本体论抽象把数据库结构“升华”为业务语义知识图谱：AI 从表名、字段与注释中识别概念（实体）'
              '与关系，构建实体间的语义关联网络，并主动发现命名冲突、口径不一致等治理风险。'
              '界面提供概念数、关系数、风险数统计，支持 2D/3D 图谱切换、AI 智能提取、关系扫描与演示场景。'
              '该模块默认隐藏，可在“设置 → 标签页设置”中开启。')
    image(doc, '03-ontology.png', '图 11　本体论抽象（知识图谱）')

    toc_anchor(doc, '七、数据血缘', size=SZ_TITLE - 6)
    body(doc, '数据血缘用于回答“这张表的数据从哪来、又被谁引用”。选择数据库后点击“分析 / 刷新”，'
              '平台会基于外键约束生成血缘关系图（箭头从被引用表指向引用表，即数据流向），'
              '并在下方列出外键明细。评估表结构变更影响范围时，先看血缘是最稳妥的做法。'
              '该模块默认隐藏，可在“设置 → 标签页设置”中开启。')
    image(doc, '04-lineage.png', '图 12　数据血缘分析')

    toc_anchor(doc, '八、接口分发', size=SZ_TITLE - 6)
    body(doc, '接口分发把平台的数据库与治理能力以接口形式对外提供。界面左侧为“平台 / 接口”两级清单，'
              '右侧为接口详情：接口名称、启用状态、接口路径、请求方法、接口类型、关联数据库、转发 URL、'
              'SQL 语句、参数说明与调用示例，并提供“🧪 测试”做在线调试。'
              '接口 Key 在“设置 → API Key”中管理，与 Agent 服务共用。')
    image(doc, '12-api.png', '图 13　接口分发与接口详情')

    toc_anchor(doc, '九、Agent 服务（MCP）', size=SZ_TITLE - 6)
    body(doc, 'Agent 服务通过 MCP（Model Context Protocol）协议把数据工具箱接入各种智能体客户端'
              '（Cursor、Claude Desktop、Cherry Studio、Dify 等），让智能体可以直接查询数据库、'
              '读取表结构、调用平台接口。页面提供：启用开关、连接信息（服务地址 / API Key / MCP 端口）、'
              '按客户端一键生成配置片段、以及安全配置（如只读模式，仅允许执行查询类语句）。')
    body(doc, '使用步骤：启用 MCP → 生成并复制 API Key → 选择客户端复制配置 → 在客户端中验证连通性。'
              'MCP 端口留空或填 0 表示复用主服务器端口，修改后需重启服务生效。')
    image(doc, '13-agent.png', '图 14　Agent 服务（MCP）配置')

    toc_anchor(doc, '十、智能助手', size=SZ_TITLE - 6)

    h2(doc, '（一）模型配置与自动获取模型')
    body(doc, '首次使用需在“设置 → LLM 配置”中填写 AI 服务 URL、API Key 与模型名称。'
              '为减少手误，模型名称支持“🔄 自动获取模型列表”：点击后平台自动访问上游服务的模型清单接口，'
              '把可用模型灌入下拉候选，既可以下拉选择，也可以继续手工输入。'
              '向量（Embedding）模型一栏同样支持自动获取，且 Key 留空时自动复用上面的主 Key。')
    image(doc, '19-ai-settings.png', '图 15　LLM 配置与“自动获取模型列表”')
    body(doc, '“模型管理”模块另提供模型清单维护：大模型管理用于配置 LLM、Rerank、向量化、'
              '语音识别等 AI 服务的转发接口；小模型管理用于把自定义数据处理逻辑（JavaScript）'
              '注册成可被引用的处理步骤。该模块默认隐藏，可在“设置 → 标签页设置”中开启。')

    h2(doc, '（二）能力检测与开关')
    body(doc, '配置保存后平台会探测模型能力（函数调用、流式输出、JSON 模式、思考模式等），'
              '并给出上下文窗口建议值。若上游服务对某些参数兼容性不佳，可手动关闭对应开关。')

    h2(doc, '（三）表检索（RAG）配置')
    body(doc, '当数据库中表数量较多时，全量表结构会占用大量上下文。可在设置中启用表检索，'
              '让平台在对话或生成任务前先筛选出相关表再交给模型，既省成本也更准确。'
              '检索策略支持关键词、向量、图谱与混合模式，可调整返回表数量与各策略权重。')

    h2(doc, '（四）对话使用')
    body(doc, '在“智能助手”标签页中可直接对话：输入 @ 可以引用某个数据库或模块，'
              '界面也提供“看看我有哪些数据”“这个表长什么样”“做个数据看板”“查一下最近异常”'
              '“帮我创建一个接口”等快捷入口。涉及写库或对外发送的操作，平台仍会要求人工确认。')
    image(doc, '18-ai-assistant.png', '图 16　智能助手对话界面')

    toc_anchor(doc, '十一、模型管理', size=SZ_TITLE - 6)
    body(doc, '模型管理分为“大模型”和“小模型”：大模型用于配置 LLM、Rerank、向量化、语音识别等'
              'AI 服务的转发接口；小模型用 JavaScript 实现自定义数据处理逻辑，可调用已注册的数据库，'
              '相当于把一段固定的加工逻辑封装成可复用的“处理单元”。该模块默认隐藏，可在设置中开启。')
    image(doc, '14-models.png', '图 17　模型管理')

    toc_anchor(doc, '十二、数据质量审核', size=SZ_TITLE - 6)
    body(doc, '数据质量审核面向“填报类”数据，由四个部分组成：')
    bullet(doc, '1. 规则配置（规则树）：按 NM（6 位编码）/ XH（层级编码）/ 名称 / 类别 / SQL / 参数 / AI 复核原则'
                '维护规则，支持下载模板后批量粘贴导入或从 Excel 导入，SQL 用 Oracle 方言编写并由平台自动转换方言；'
                'SQL 中用 {{占位符}} 表示待填参数，未填参数的规则会直接报“参数未配置”，避免出现难懂的语法报错。')
    bullet(doc, '2. 定时任务：按周期自动执行审核规则，并可按周期统计填报完成情况。')
    bullet(doc, '3. 执行记录：查看每次审核的执行结果与问题清单，逐级定位到具体数据项。')
    bullet(doc, '4. 报告模板：定义审核报告的呈现格式，便于导出与上报。')
    image(doc, '15-quality.png', '图 18　数据质量审核（规则配置，含“AI 复核原则”）')

    h2(doc, '（一）SQL 审核 + AI 复核：两层判定')
    body(doc, '一条规则的判定分两层。第一层是 SQL 审核：执行规则中的 SQL，查出违规行即判为“不通过”。'
              '第二层是 AI 复核（可选）：只有当规则本身填写了“AI 复核原则”时才会触发——'
              '审核不通过时，平台把这段原则、规则 SQL、执行结果一并交给大模型，'
              '判断“这次不通过是不是规则过严导致的误判”，并给出置信度、理由与规则修改建议。'
              '这样可以把“数据真有问题”和“规则写得过严”区分开，减少无谓的整改通知。')
    body(doc, 'AI 复核原则的填写与维护有两种方式：')
    bullet(doc, '· 批量导入：在导入表最后增加一列「AI复核原则」（第 7 列，可选），表头写“AI复核原则”或'
                '“人工审核规则”均可，列顺序可调；下载的导入模板已带该列说明。')
    bullet(doc, '· 单条维护：在规则编辑表单中直接填写“AI 复核原则”输入框，适用少量调整。')
    body(doc, '关于兼容性与费用，请留意：')
    bullet(doc, '· 旧的 5 列 / 6 列导入文件照常导入，缺该列时该规则不做 AI 复核（留空即等于以 SQL 结果为准）。')
    bullet(doc, '· 留空的规则完全不调用大模型，不产生调用费用。')
    bullet(doc, '· AI 结论仅供参考：报告中每条 AI 结论都标注“※ 本条须由人类专家最终校核”，不能作为最终定论。')
    body(doc, 'AI 复核是“按规则开关”的，两个执行入口都要打开对应开关才会生效：'
              '手动一键审核时勾选“启用 AI 校核”；定时任务在任务弹窗中勾选“启用 AI 校核”（默认开启）。'
              '若某次执行未启用，结果与报告中会明确提示“本次未启用 AI 校核”，不会出现看似复核过、实为旧结论的情况。')
    image(doc, '22-qa-ai-review.png', '图 19　审核结果中的“SQL 审核结果”与“AI 复核结论”')

    h2(doc, '（二）定时任务的审核项与开关')
    body(doc, '定时任务弹窗中的审核项规则树只保留一列「执行」：勾选即代表该规则参与本次审核。'
              '某条规则是否送 AI 复核，由“规则自带的复核原则 + 任务级的启用 AI 校核”共同决定，'
              '不需要在任务里逐条重复勾选，规则调整后任务自动跟随。')
    image(doc, '21-qa-schedule.png', '图 20　定时任务：审核项单列勾选与“启用 AI 校核”')

    toc_anchor(doc, '十三、应用广场', size=SZ_TITLE - 6)
    body(doc, '应用广场用于创建与管理自定义应用，支持纯前端的 HTML + CSS + JS 应用：'
              '新建应用后可在线编辑页面代码、绑定数据、预览效果，并可将应用设为公开访问，'
              '把地址分享给同事或嵌入其它系统。平台同时提供预制看板模板，可作为起步脚手架。')
    image(doc, '16-apps.png', '图 21　应用广场')

    toc_anchor(doc, '十四、大屏编辑器', size=SZ_TITLE - 6)
    body(doc, '大屏编辑器用于拼装数据大屏：从组件库中拖入图表、表格、指标卡等组件，绑定数据源后'
              '调整布局与样式，支持保存、预览与分享。组件与模板均在本地打包，无需联网即可使用。'
              '该模块默认隐藏，可在“设置 → 标签页设置”中开启。')
    image(doc, '17-screen-editor.png', '图 22　大屏编辑器')


def build_faq_and_appendix(doc):
    toc_anchor(doc, '十五、常见问题与故障排查', size=SZ_TITLE - 6)
    body(doc, '1. 运行任务时提示“请上传文件”？')
    body(doc, '　　说明任务需要输入。请在“数据输入”区域选择文件或粘贴文本后再运行；'
              '合并执行的任务需要把全部相关文件一次性选中。', indent=0)
    body(doc, '2. 任务执行成功但输出文件为空/格式不符？')
    body(doc, '　　先看执行日志中的解析条数是否符合预期。若输入文件结构与示例不同，'
              '需要调整脚本中的解析规则；也可用“AI 新建”基于现有任务生成一版再微调。', indent=0)
    body(doc, '3. AI 相关功能报错？')
    body(doc, '　　在“设置 → LLM 配置”中核对 URL、Key 与模型名；'
              '点击“自动获取模型列表”可快速判断是网络问题还是配置问题。模型名写错是最常见原因。', indent=0)
    body(doc, '4. 升级后界面没有变化？')
    body(doc, '　　浏览器缓存了旧版静态文件。按 Ctrl+F5 强制刷新；'
              '安装新版本时请确保 index.html 与 js/、css/ 目录同时更新。', indent=0)
    body(doc, '5. 分享链接是否需要收回？')
    body(doc, '　　需要。分享链接不要求登录，用完请在任务详情页关闭分享。', indent=0)
    body(doc, '6. 某条规则为什么没有 AI 复核结论？')
    body(doc, '　　两种常见原因：该规则没有填写“AI 复核原则”（留空即以 SQL 审核结果为准，不调用 AI）；'
              '或本次执行没有勾选“启用 AI 校核”。报告与结果区会写明未启用的原因，便于核对。', indent=0)
    body(doc, '7. AI 判断为“疑似规则过严误判”，可以直接不改数据吗？')
    body(doc, '　　不可以。AI 结论仅供参考，须由业务/数据专家结合规则口径复核后决定：'
              '确属误判就修改规则（如增加可空字段白名单、放开阈值），确属数据问题则照常整改。', indent=0)
    body(doc, '8. 如何备份数据？')
    body(doc, '　　停止服务后备份 data/ 目录（含 data-store.json）与 server.config.json 即可；'
              '恢复时覆盖回原位置再启动服务。', indent=0)

    toc_anchor(doc, '附录 A　治理任务内置 API（gov.*）速查')
    body(doc, '下表为治理任务脚本中可直接使用的主要接口，更完整的说明见平台内“接口分发 → API 参考”。')
    table(doc, [
        ['接口', '用途'],
        ['gov.log(msg)', '向执行日志输出一行信息'],
        ['gov.showTable(data)', '以表格形式预览结构化结果'],
        ['gov.readExcel(file)', '读取 Excel，返回工作簿对象'],
        ['gov.readCSV(text)', '解析 CSV 文本为二维数组'],
        ['gov.readWord(file)', '读取 Word 正文，返回 { value }'],
        ['gov.readWordTables(file)', '读取 Word 中的表格（含列宽、样式），用于套用模板'],
        ['gov.word()', 'Word 构建器：heading / paragraph / table / tableFromTemplate / save'],
        ['gov.writeExcel(name, data, opts?)', '写出 Excel（支持单元格样式）'],
        ['gov.writeCSV(name, data)', '写出 CSV 文件'],
        ['gov.writeText(name, content)', '写出纯文本文件'],
        ['gov.writeJSON(name, data)', '写出 JSON 文件'],
        ['gov.fillWordTemplate(tpl, data, out)', '按模板填充 Word 并输出'],
        ['gov.fillExcelTemplate(tpl, data, out)', '按模板填充 Excel 并输出'],
        ['gov.parseFilename(name, opts?)', '从文件名解析单位与日期'],
        ['gov.parseWordStructure(file, opts?)', '解析 Word 的标题层级结构'],
        ['gov.callAI(prompt)', '调用已配置的大模型，返回文本结果'],
        ['gov.getDbType()', '获取当前关联数据库类型'],
        ['gov.getDatabases()', '列出可用数据库'],
        ['gov.querySQL(sql, params?)', '在当前数据库执行查询，返回行数组'],
        ['gov.executeSQL(sql, params?)', '在当前数据库执行写入，返回影响行数'],
        ['gov.querySQLForDb(dbId, sql, params?)', '在指定数据库执行查询'],
        ['gov.executeSQLForDb(dbId, sql, params?)', '在指定数据库执行写入'],
    ], widths=[6.2, 9.3], size=9)

    toc_anchor(doc, '附录 B　内置示例任务速查')
    table(doc, [
        ['示例任务', '输入', '输出 / 用途'],
        ['产品汇总Word生成（套用表格模板）', '表格模板 Word + 产品介绍 Word', '套用模板样式的产品汇总 Word（每个产品一张表）'],
        ['公文Word转Excel（正则抽取）', '公文类 Word', '按正则规则抽取要素，输出 Excel'],
        ['公文Word转Excel（AI抽取）', '公文类 Word', '由 AI 理解正文后抽取要素，输出 Excel'],
        ['Word文档内容提取', 'Word 文件', '提取正文与结构，便于后续加工'],
        ['Excel数据解析入库', 'Excel 文件', '解析并写入指定数据库表'],
        ['CSV文本解析', 'CSV 文件或文本', '解析为结构化数据'],
        ['国际新闻入库', '新闻文本/文件', '整理后入库'],
        ['综合日报生成器', '多个日报 Word', '合并生成综合日报（合并执行模式）'],
        ['数据完整性检查', '—（定时）', '定时巡检数据完整性并输出结论'],
        ['数据库表行数统计', '—（定时）', '定时统计表行数，用于监控数据增长'],
        ['测试任务', '—', '用于验证环境与脚本的最小示例'],
    ], widths=[5.2, 4.3, 6.0], size=9)
    body(doc, '提示：示例任务可直接“运行”查看效果，也可用“📄 复制”或“✨ AI 新建”派生出符合自身业务的新任务。')

    toc_anchor(doc, '附录 C　本说明书的更新维护')
    body(doc, '本说明书随版本发布一并打包，目录下的“使用说明书.docx”即为当前版本。'
              '为确保说明与系统功能始终一致，更新约定如下：')
    bullet(doc, '1. 新增或调整功能时，同步修改 docs/manual/build_manual.py 中的对应章节文字与表格。')
    bullet(doc, '2. 界面发生变化时，重新运行 docs/manual/take_screenshots.js 抓取最新截图，再重新生成说明书。')
    bullet(doc, '3. 生成命令：python3 docs/manual/build_manual.py，输出到 docs/使用说明书.docx。')
    bullet(doc, '4. 发布流水线会在打包时自动重新生成说明书，并放入安装包根目录。')
    bullet(doc, '5. 每次发布前确认：版本号、示例任务清单、界面截图三处与系统实际一致。')


TOC_PAGES = {}


def _toc_line(doc, text, page):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(24)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    set_indent_chars(par, first=0)
    pPr = par._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:leader'), 'dot')
    tab.set(qn('w:pos'), '8930')
    tabs.append(tab)
    pPr.append(tabs)
    run = par.add_run(text + '\t' + (str(page) if page else ''))
    set_font(run, east=F_FS, size=SZ_BODY)
    return par


def build_toc(doc, items, pages):
    body(doc, '本说明书共分十五章及三个附录，主要内容目录如下：', indent=2, space_after=8)
    for text in items:
        _toc_line(doc, text, pages.get(text, ''))


def build_document(toc_items=None, pages=None):
    global TOC_ANCHORS
    TOC_ANCHORS = []
    doc = setup_document()
    build_cover(doc, toc_items, pages)
    build_intro(doc)
    doc.add_page_break()
    build_install(doc)
    doc.add_page_break()
    build_login_and_nav(doc)
    doc.add_page_break()
    build_database(doc)
    doc.add_page_break()
    build_governance(doc)
    doc.add_page_break()
    build_other_modules(doc)
    doc.add_page_break()
    build_faq_and_appendix(doc)
    return doc


def extract_page_map(pdf_path, anchors):
    """从 PDF 里找出每个标题所在页码（跳过封面与目录页）"""
    import re as _re
    import subprocess
    try:
        out = subprocess.run(['pdftotext', '-layout', pdf_path, '-'],
                             capture_output=True, text=True, timeout=180).stdout
    except Exception:
        return {}
    if not out:
        return {}
    pages = out.split('\f')
    norm = [_re.sub(r'\s+', '', p) for p in pages]
    # 跳过封面与目录页：用目录专属标记定位（正文里也有“目录结构”“目录”等词，不能全局搜）
    start = 0
    for i, pg in enumerate(norm):
        if '主要内容目录如下' in pg:
            start = i + 1
    if start == 0:
        for i in range(min(6, len(norm))):
            if '目录' in norm[i]:
                start = i + 1
                break
    page_map = {}
    for a in anchors:
        key = _re.sub(r'\s+', '', a)
        if not key:
            continue
        for i in range(start, len(norm)):
            if key in norm[i]:
                page_map[a] = i + 1
                break
    return page_map


def main():
    import shutil
    import subprocess
    import tempfile

    # 第一遍：仅为收集目录条目
    build_document(None, None)
    anchors = list(TOC_ANCHORS)

    # 第二遍：目录条目一致但页码留空（与最终排版一致），用于测定页码
    tmpdir = tempfile.mkdtemp(prefix='dtmanual-')
    tmp_docx = os.path.join(tmpdir, 'manual.docx')
    build_document(anchors, {}).save(tmp_docx)

    pages = {}
    if shutil.which('soffice'):
        try:
            subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmpdir, tmp_docx],
                           capture_output=True, timeout=420)
            pdf = os.path.join(tmpdir, 'manual.pdf')
            if os.path.exists(pdf):
                pages = extract_page_map(pdf, anchors)
                print('已测定页码：%d/%d 个标题' % (len(pages), len(anchors)))
        except Exception as e:  # noqa: BLE001
            print('页码测定跳过：%s' % e)
    else:
        print('未找到 soffice，目录将不含页码')

    # 第三遍：写入页码
    TOC_PAGES.clear()
    TOC_PAGES.update(pages)
    doc2 = build_document(anchors, pages)
    if list(TOC_ANCHORS) != anchors:
        print('警告：两遍的目录条目不一致，页码可能不准')
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc2.save(OUT)
    print('已生成：%s（%.1f KB）' % (OUT, os.path.getsize(OUT) / 1024.0))


if __name__ == '__main__':
    main()
